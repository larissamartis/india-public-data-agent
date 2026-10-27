from datetime import date, datetime
from io import BytesIO

import duckdb
import httpx
import openpyxl
import pytest

from bharat.ingest import gst
from bharat.ingest.common import make_client


def build_workbook(fy="2025-26") -> bytes:
    """A small copy of the real layout: title rows, merged period labels, two-row header.

    Months are January (as text) and February (as a date cell) of the FY's second year.
    """
    end_year = int(fy[:4]) + 1
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Collections-Statewise"
    ws.append(["State Wise Domestic GST Collection"])
    ws.append(["Data upto: 31st Mar 2026 (Rs. in crore)"])
    ws.append([])
    # Row 4: period labels, each merged over four cells. The FY total comes first.
    ws.append(
        ["State CD", "State", fy, None, None, None, f"Jan-{end_year % 100:02d}", None, None, None]
        + [datetime(end_year, 2, 1), None, None, None]
    )
    ws.merge_cells("C4:F4")
    ws.merge_cells("G4:J4")
    ws.merge_cells("K4:N4")
    ws.append([None, None] + ["CGST", "SGST", "IGST", "TOTAL"] * 3)
    ws.append([27, "Maharashtra", 100, 120, 130, 350, 50, 60, 65, 175, 50, 60, 65, 175])
    ws.append([32, "Kerala", 10, 12, "-", 22, 5, 6, None, 11, 5, 6, None, 11])
    ws.append([97, "Other Territory", 2, 2, 2, 6, 1, 1, 1, 3, 1, 1, 1, 3])
    ws.append([99, "OIDAR", 0, 0, 40, 40, 0, 0, 20, 20, 0, 0, 20, 20])
    ws.append([])
    ws.append(["Note: figures are provisional and may vary on finalisation."])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    ("label", "expected"),
    [
        ("Apr-25", date(2025, 4, 1)),
        ("April-2025", date(2025, 4, 1)),
        ("Jan 2026", date(2026, 1, 1)),
        (datetime(2025, 11, 1), date(2025, 11, 1)),
        ("March", date(2026, 3, 1)),
        ("May", date(2025, 5, 1)),
        ("2025-26", None),
        ("Total", None),
        (None, None),
    ],
)
def test_parse_month(label, expected):
    assert gst.parse_month(label, "2025-26") == expected


def test_fiscal_years_follow_april_to_march():
    assert gst.fiscal_years(date(2026, 3, 31))[-1] == "2025-26"
    assert gst.fiscal_years(date(2026, 4, 1))[-1] == "2026-27"
    assert gst.fiscal_years(date(2026, 4, 1))[0] == "2017-18"


def test_parse_workbook_unpivots_months_and_splits_non_states():
    states, others = gst.parse_workbook(build_workbook(), "2025-26")
    assert {(r["state_name"], r["month"]) for r in states} == {
        ("Maharashtra", "2026-01-01"),
        ("Maharashtra", "2026-02-01"),
        ("Kerala", "2026-01-01"),
        ("Kerala", "2026-02-01"),
    }
    kerala_jan = next(
        r for r in states if r["state_name"] == "Kerala" and r["month"] == "2026-01-01"
    )
    assert kerala_jan == {
        "fy": "2025-26",
        "month": "2026-01-01",
        "state_code_src": "32",
        "state_name": "Kerala",
        "cgst": 5.0,
        "sgst": 6.0,
        "igst": None,
        "total": 11.0,
    }
    assert {r["state_name"] for r in others} == {"Other Territory", "OIDAR"}


def test_layout_change_fails_loudly():
    wb = openpyxl.Workbook()
    wb.active.title = "Something else"
    out = BytesIO()
    wb.save(out)
    with pytest.raises(gst.GstFormatError, match="Collections-Statewise"):
        gst.parse_workbook(out.getvalue(), "2025-26")


def test_run_downloads_each_year_and_writes_bronze(tmp_path):
    requested = []

    def handler(request):
        requested.append(str(request.url))
        if "2026-27" in str(request.url):
            return httpx.Response(404)
        fy = str(request.url).rsplit("_", 1)[1].removesuffix(".xlsx")
        return httpx.Response(200, content=build_workbook(fy))

    client = make_client(transport=httpx.MockTransport(handler))
    paths = gst.run(tmp_path, client=client, today=date(2026, 4, 10), sleep=lambda s: None)

    assert len(requested) == 10  # 2017-18 to 2026-27; the 2026-27 file isn't out yet
    assert len(paths) == 2 and all(p.exists() for p in paths)
    rows = duckdb.sql(
        f"select count(*), count(distinct state_name) from '{paths[0].as_posix()}'"
    ).fetchall()[0]
    assert rows == (9 * 2 * 2, 2)  # 9 years x 2 states x 2 months
    assert (tmp_path / "raw" / "gst" / "statewise_GST_collection_2025-26.xlsx").exists()
