"""GST collections by state and month, from the GST statistics page.

One workbook per financial year. The `Collections-Statewise` sheet has a two-row header: a
row of period labels (merged across four cells) above a row of CGST, SGST, IGST and TOTAL.
The parser finds those rows instead of assuming fixed positions, skips the full-year block,
and fails with a clear message if the layout changes.

Figures are provisional, in ₹ crore, and count tax by the state where the taxpayer is
registered, not where goods are consumed.
"""

from __future__ import annotations

import calendar
import logging
import time
from collections.abc import Callable
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import httpx
import openpyxl

from bharat.contracts import load_source_contract
from bharat.ingest.bronze import save_bronze
from bharat.ingest.common import make_client, request_with_retry, sha256_bytes, to_float
from bharat.reference.states import normalise_name

log = logging.getLogger(__name__)

PAGE_URL = "https://www.gst.gov.in/download/gststatistics"
FILE_URL = (
    "https://tutorial.gst.gov.in/offlineutilities/gst_statistics/statewise_GST_collection_{fy}.xlsx"
)
SHEET = "Collections-Statewise"
COMPONENTS = ("CGST", "SGST", "IGST", "TOTAL")
FIRST_FY_START = 2017
# Rows that are not states: kept in their own table so state totals stay clean.
NON_STATE_ROWS = {
    "other territory",
    "oidar",
    "total",
    "grand total",
    "centre jurisdiction",
    "others",
}
_NAME_HEADERS = {"state", "state name", "state ut", "states ut", "name of the state", "state uts"}
_CODE_HEADERS = {"state cd", "state code", "code", "state cd code"}
_MONTHS = {name.lower(): i for i, name in enumerate(calendar.month_abbr) if name}

STATE_COLUMNS = {
    "fy": "VARCHAR",
    "month": "DATE",
    "state_code_src": "VARCHAR",
    "state_name": "VARCHAR",
    "cgst": "DOUBLE",
    "sgst": "DOUBLE",
    "igst": "DOUBLE",
    "total": "DOUBLE",
}


class GstFormatError(ValueError):
    """The workbook no longer has the layout this parser expects."""


def fiscal_years(today: date, first_start: int = FIRST_FY_START) -> list[str]:
    """FY labels from 2017-18 up to the current FY (April to March)."""
    current_start = today.year if today.month >= 4 else today.year - 1
    return [f"{y}-{(y + 1) % 100:02d}" for y in range(first_start, current_start + 1)]


def parse_month(label: object, fy: str) -> date | None:
    """Turn a header label into the first day of its month; None for non-month labels."""
    if isinstance(label, datetime | date):
        return date(label.year, label.month, 1)
    text = "" if label is None else str(label).strip()
    if not text:
        return None
    for fmt in ("%b-%y", "%b-%Y", "%B-%y", "%B-%Y", "%b %Y", "%B %Y", "%b %y", "%b'%y", "%b, %Y"):
        try:
            parsed = datetime.strptime(text, fmt)
        except ValueError:
            continue
        return date(parsed.year, parsed.month, 1)
    if text.isalpha() and text[:3].lower() in _MONTHS and len(text) <= 9:
        month = _MONTHS[text[:3].lower()]
        start = int(fy[:4])
        return date(start if month >= 4 else start + 1, month, 1)
    return None


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _cell(row: list, col: int | None) -> object:
    if col is None or col >= len(row):
        return None
    return row[col]


def _code(value: object) -> str | None:
    if value is None or _text(value) == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def parse_workbook(content: bytes, fy: str) -> tuple[list[dict], list[dict]]:
    """Return (state rows, other rows), one row per name per month."""
    workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if SHEET not in workbook.sheetnames:
            raise GstFormatError(f"sheet {SHEET!r} not found; sheets are {workbook.sheetnames}")
        grid = [list(row) for row in workbook[SHEET].iter_rows(values_only=True)]
    finally:
        workbook.close()

    comp_row = next(
        (i for i, row in enumerate(grid[:30]) if any(_text(v).upper() == "CGST" for v in row)),
        None,
    )
    if comp_row is None:
        raise GstFormatError(f"{fy}: no CGST header in the first 30 rows")
    comps = {
        c: _text(v).upper() for c, v in enumerate(grid[comp_row]) if _text(v).upper() in COMPONENTS
    }
    first_comp, last_comp = min(comps), max(comps)

    period_row = next(
        (
            i
            for i in range(comp_row - 1, max(comp_row - 4, -1), -1)
            if any(_text(v) for v in grid[i][first_comp : last_comp + 1])
        ),
        None,
    )
    if period_row is None:
        raise GstFormatError(f"{fy}: no period labels above the CGST row")

    # Merged header cells keep their label only in the first cell, so carry it to the right.
    labels: dict[int, object] = {}
    current: object = None
    for c in range(first_comp, last_comp + 1):
        value = _cell(grid[period_row], c)
        if _text(value):
            current = value
        labels[c] = current

    name_col = code_col = None
    for i in range(period_row, comp_row + 1):
        for c, value in enumerate(grid[i][:first_comp]):
            header = normalise_name(_text(value))
            if header in _CODE_HEADERS:
                code_col = c
            elif header in _NAME_HEADERS:
                name_col = c
    if name_col is None:
        name_col = first_comp - 1
        log.warning("%s: no 'State' header found; using column %s for names", fy, name_col)

    by_month: dict[date, dict[str, int]] = {}
    for c, comp in comps.items():
        month = parse_month(labels.get(c), fy)
        if month is not None:
            by_month.setdefault(month, {})[comp] = c
    if not by_month:
        seen = sorted({_text(v) for v in labels.values() if _text(v)})
        raise GstFormatError(f"{fy}: no month columns found; header labels were {seen}")

    value_cols = [c for cols in by_month.values() for c in cols.values()]
    states: list[dict] = []
    others: list[dict] = []
    for row in grid[comp_row + 1 :]:
        name = _text(_cell(row, name_col))
        if not name or all(to_float(_cell(row, c)) is None for c in value_cols):
            continue  # blank lines and notes have no figures
        target = others if normalise_name(name) in NON_STATE_ROWS else states
        for month, cols in sorted(by_month.items()):
            record = {
                "fy": fy,
                "month": month.isoformat(),
                "state_code_src": _code(_cell(row, code_col)),
                "state_name": name,
            }
            for comp in COMPONENTS:
                col = cols.get(comp)
                record[comp.lower()] = to_float(_cell(row, col)) if col is not None else None
            target.append(record)
    return states, others


def run(
    out_dir: Path,
    *,
    client: httpx.Client | None = None,
    today: date | None = None,
    delay: float = 2.0,
    sleep: Callable[[float], None] = time.sleep,
) -> list[Path]:
    """Download every FY workbook, parse it and write the bronze tables."""
    today = today or date.today()
    client = client or make_client()
    raw_dir = out_dir / "raw" / "gst"
    raw_dir.mkdir(parents=True, exist_ok=True)
    years = fiscal_years(today)
    states: list[dict] = []
    others: list[dict] = []
    for i, fy in enumerate(years):
        url = FILE_URL.format(fy=fy)
        response = request_with_retry(client, "GET", url, sleep=sleep)
        if response.status_code == 404 and fy == years[-1]:
            log.warning("no file yet for the current year %s", fy)
            continue
        response.raise_for_status()
        (raw_dir / f"statewise_GST_collection_{fy}.xlsx").write_bytes(response.content)
        ref = f"{url}#sha256={sha256_bytes(response.content)}"
        year_states, year_others = parse_workbook(response.content, fy)
        log.info("%s: %s state rows, %s other rows", fy, len(year_states), len(year_others))
        for record in (*year_states, *year_others):
            record["_source_ref"] = ref
        states.extend(year_states)
        others.extend(year_others)
        if i < len(years) - 1:
            sleep(delay)

    paths = [
        save_bronze(
            states, STATE_COLUMNS, load_source_contract("gst_state_monthly"), out_dir, today=today
        )
    ]
    if others:
        paths.append(
            save_bronze(
                others,
                STATE_COLUMNS,
                load_source_contract("gst_other_monthly"),
                out_dir,
                today=today,
            )
        )
    else:
        log.warning("no non-state rows (Other Territory, OIDAR) found in any year")
    return paths
