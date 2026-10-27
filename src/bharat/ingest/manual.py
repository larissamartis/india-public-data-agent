"""Hand-downloaded files: fetch them from the `raw-inputs` release, check them, preview them.

Files are listed in `data/manifest.yml`. Before any parsing, every file for a source is
downloaded from the release and its SHA-256 compared with the manifest. The preview
("inspect") shows the first rows of each file, so a parser can be written against the real
layout rather than a guess.
"""

from __future__ import annotations

import csv
import os
from io import BytesIO, StringIO
from pathlib import Path

import httpx
import openpyxl

from bharat.ingest.common import make_client, request_with_retry
from bharat.ingest.manifest import Manifest, ManifestFile, load_manifest, verify_files

MANIFEST_PATH = Path("data/manifest.yml")
DEFAULT_REPO = "larissamartis/india-public-data-agent"
RELEASE_URL = "https://github.com/{repo}/releases/download/raw-inputs/{name}"
PREVIEW_ROWS = 15
PREVIEW_WIDTH = 40
PREVIEW_CHARS = 6000


class ManualFileError(RuntimeError):
    pass


def entries_for(source: str, manifest_path: Path = MANIFEST_PATH) -> list[ManifestFile]:
    entries = load_manifest(manifest_path).for_source(source)
    if not entries:
        raise ManualFileError(
            f"no files for '{source}' in {manifest_path}; upload them to the raw-inputs "
            "release and list them in the manifest first"
        )
    return entries


def fetch_files(
    entries: list[ManifestFile],
    folder: Path,
    *,
    client: httpx.Client | None = None,
    repo: str | None = None,
) -> list[Path]:
    """Download any listed file not already in `folder`, then check every fingerprint."""
    client = client or make_client()
    repo = repo or os.environ.get("GITHUB_REPOSITORY") or DEFAULT_REPO
    folder.mkdir(parents=True, exist_ok=True)
    for entry in entries:
        path = folder / entry.name
        if path.exists():
            continue
        response = request_with_retry(client, "GET", RELEASE_URL.format(repo=repo, name=entry.name))
        if response.status_code == 404:
            raise ManualFileError(f"{entry.name} is listed but not on the raw-inputs release")
        response.raise_for_status()
        path.write_bytes(response.content)
    problems = verify_files(Manifest(files=entries), folder)
    if problems:
        raise ManualFileError(
            "manual files failed the manifest check:\n  - " + "\n  - ".join(problems)
        )
    return [folder / entry.name for entry in entries]


def _short(value: object) -> str:
    text = "" if value is None else " ".join(str(value).split())
    return text if len(text) <= PREVIEW_WIDTH else text[: PREVIEW_WIDTH - 1] + "…"


def _table_preview(rows) -> list[str]:
    lines = []
    for row in rows:
        cells = [_short(v) for v in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            lines.append(" | ".join(cells))
        if len(lines) == PREVIEW_ROWS:
            break
    return lines


def _xlsx_preview(content: bytes) -> list[str]:
    book = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)
    lines = []
    for sheet in book.worksheets:
        lines.append(f"[sheet {sheet.title}: {sheet.max_row} rows x {sheet.max_column} columns]")
        lines.extend(_table_preview(sheet.iter_rows(values_only=True)))
    return lines


def _xls_preview(content: bytes) -> list[str]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    lines = []
    for sheet in book.sheets():
        lines.append(f"[sheet {sheet.name}: {sheet.nrows} rows x {sheet.ncols} columns]")
        lines.extend(_table_preview(sheet.row_values(i) for i in range(sheet.nrows)))
    return lines


def _csv_preview(content: bytes) -> list[str]:
    text = content.decode("utf-8-sig", errors="replace")
    return _table_preview(csv.reader(StringIO(text)))


def _pdf_preview(content: bytes) -> list[str]:
    import pdfplumber

    lines = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        lines.append(f"[pdf: {len(pdf.pages)} pages]")
        for number, page in enumerate(pdf.pages[:2], start=1):
            lines.append(f"[page {number}]")
            text_lines = [line for line in (page.extract_text() or "").splitlines() if line.strip()]
            lines.extend(_short(line) for line in text_lines[: PREVIEW_ROWS * 2])
    return lines


PREVIEWERS = {
    ".xlsx": _xlsx_preview,
    ".xls": _xls_preview,
    ".csv": _csv_preview,
    ".pdf": _pdf_preview,
}


def inspect_file(path: Path) -> str:
    """A short, plain-text look at the start of a file, for writing its parser."""
    content = path.read_bytes()
    previewer = PREVIEWERS.get(path.suffix.lower())
    header = f"{path.name} ({len(content):,} bytes)"
    if previewer is None:
        return f"{header}: no preview for this file type"
    try:
        lines = previewer(content)
    except Exception as error:  # a preview must never stop the other files being shown
        return f"{header}: could not read ({type(error).__name__}: {error})"
    text = "\n".join([header, *lines])
    return text if len(text) <= PREVIEW_CHARS else text[: PREVIEW_CHARS - 1] + "…"
